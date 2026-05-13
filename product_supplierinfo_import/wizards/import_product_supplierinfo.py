# © 2025 Seges
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html)

import logging
import binascii
import tempfile

from openpyxl import load_workbook

from odoo import models, fields, api, _
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)


class ImportProduct_supplierinfo(models.TransientModel):
    _name = 'import.product_supplierinfo'
    _description = _('ImportProduct_supplierinfo')
    
    file = fields.Binary(string="Archivo", required=True,
                         help="Archivo a subir")

    def action_product_supplierinfo_import(self):
        
        try:
            try:
                file_pointer = tempfile.NamedTemporaryFile(delete=False,
                                                            suffix=".xlsx")
                file_pointer.write(binascii.a2b_base64(self.file))
                file_pointer.seek(0)
                workbook = load_workbook(filename=file_pointer.name)
                sheet = workbook['Sheet1']
                max_rows = sheet.max_row
                if max_rows >2000:
                    raise UserError(_("Archivo con demasiadas filas."))
            except:
                raise UserError(_("Archivo no válido"))
            
            codes = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index == 0:
                    continue
                
                if row:
                    line = index + 1
                    row_vals = row
                    
                    supplier = self.env['res.partner'].with_context(active_test=False).search(
                            [('name', '=like', row_vals[0]),('supplier_rank', '>', 0)])
                    
                    if not supplier and row_vals[0] != None:
                        raise UserError(_("El proveedor en la fila %s no existe.", str(line)))
                    
                    if row_vals[1] == None:
                        raise UserError(_("La fila %s no tiene nombre de producto.", str(line)))
                    
                    if row_vals[2] == None:
                        raise UserError(_("La fila %s no tiene referencia interna.", str(line)))
                    
                    product_template = self.env['product.template'].with_context(active_test=False).search([('default_code', '=like', row_vals[2])])
                    
                    if (len(product_template) == 1):
                        product_supplierinfo = self.env['product.supplierinfo'].with_context(active_test=False).search([
                            ('partner_id', '=', supplier.id),
                            ('product_tmpl_id', '=', product_template.id)
                        ])

                        if product_supplierinfo:
                            product_supplierinfo.unlink()

                        list_price = product_template.list_price

                        standard_price = row_vals[4]
                        if standard_price > 0.0:
                            if row_vals[5] > 0.0:
                                standard_price *= (1 - row_vals[5] / 100.0)
                            if row_vals[6] > 0.0:
                                standard_price *= (1 - row_vals[6] / 100.0)
                            if row_vals[7] > 0.0:
                                standard_price *= (1 - row_vals[7] / 100.0)

                        if row_vals[8] > 0.0:
                            list_price = standard_price / (1-(row_vals[8]/100))
                        
                        product_template.write({
                            'list_price' : list_price,
                            'standard_price' : standard_price
                        })

                        new_product_supplierinfo = self.env['product.supplierinfo'].create({
                            'partner_id' : supplier.id,
                            'product_tmpl_id' : product_template.id,
                            'product_name' : row_vals[1],
                            'product_code' : row_vals[2],
                            'min_qty' : row_vals[3],
                            'price' : row_vals[4],
                            'discount1' : row_vals[5],
                            'discount2' : row_vals[6],
                            'discount3' : row_vals[7],
                            'benefit' : row_vals[8],
                            'date_start' : row_vals[9],
                            'date_end' : row_vals[10],
                        })
                    elif len(product_template) == 0:
                        list_price = 0.0

                        standard_price = row_vals[4]
                        if standard_price > 0.0:
                            if row_vals[5] > 0.0:
                                standard_price *= (1 - row_vals[5] / 100.0)
                            if row_vals[6] > 0.0:
                                standard_price *= (1 - row_vals[6] / 100.0)
                            if row_vals[7] > 0.0:
                                standard_price *= (1 - row_vals[7] / 100.0)

                        if row_vals[8] > 0.0:
                            list_price = standard_price / (1-(row_vals[8]/100))
                        
                        product = self.env['product.template'].create({
                            'name' : row_vals[1],
                            'default_code' : row_vals[2],
                            'list_price' : list_price,
                            'standard_price' : standard_price
                        })

                        # Escribir el barcode. En este punto la variante ya existe.
                        if product:
                            product.barcode = row_vals[2]
                        
                        new_product_supplierinfo = self.env['product.supplierinfo'].create({
                            'partner_id' : supplier.id,
                            'product_tmpl_id' : product.id,
                            'product_name' : row_vals[1],
                            'product_code' : row_vals[2],
                            'min_qty' : row_vals[3],
                            'price' : row_vals[4],
                            'discount1' : row_vals[5],
                            'discount2' : row_vals[6],
                            'discount3' : row_vals[7],
                            'benefit' : row_vals[8],
                            'date_start' : row_vals[9],
                            'date_end' : row_vals[10],
                        })

                    elif len(product_template) > 1:
                        codes.append(row_vals[2])

            if codes:
                raise UserError(_("Las siguientes referencias están duplicadas: "+ " ; ".join(codes)  +"."))


        except UserError as e:
            raise UserError(str(e))
              